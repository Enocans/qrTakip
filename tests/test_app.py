import json
import os
import tempfile
import unittest
import uuid
from io import BytesIO
from unittest.mock import patch, MagicMock

from openpyxl import load_workbook
from app import app, SUBJECTS, send_whatsapp, db
from urllib.parse import parse_qs


class AppTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        app.config.update(TESTING=True, DATABASE_URL=None, DATABASE=os.path.join(self.directory.name, 'test.sqlite3'), SESSION_COOKIE_SECURE=False)
        self.client = app.test_client()
        self.csrf = self.client.get('/api/session').json['csrf']
        self.env = patch.dict(os.environ, {'TWILIO_ACCOUNT_SID':'', 'TEACHER_PHONE':''})
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.directory.cleanup()

    def data(self):
        result = dict(id=str(uuid.uuid4()), student_name='Test Öğrenci', date='2026-01-01',
                      father_phone='05321234567', teacher_phone='+905331234567', consent=True)
        for key, _, _ in SUBJECTS:
            result[key+'_correct'] = '50'
            result[key+'_wrong'] = '3'
            result[key+'_blank'] = '2'
        return result

    def post(self, data):
        return self.client.post('/api/records', json=data, headers={'X-CSRF-Token':self.csrf})

    def test_daily_total_and_idempotency(self):
        data = self.data()
        with patch('app.send_whatsapp', return_value='accepted') as send:
            response = self.post(data)
            self.assertEqual(response.status_code,201)
            self.assertEqual(response.json['total_questions'],330)
            self.assertEqual(response.json['blank'],12)
            self.assertNotIn('exam_name',response.json)
            self.assertEqual(response.json['father_phone'],'+905321234567')
            self.assertEqual(self.post(data).status_code,200)
            self.assertEqual(send.call_count,2)
        self.assertEqual(len(self.client.get('/api/records').json),1)

    def test_invalid_scores_and_consent(self):
        for key,value in [('turkce_correct','10000'),('fen_wrong','-1'),('din_wrong','1.5'),('fen_blank',True),('consent',False),('father_phone','invalid')]:
            data=self.data();data[key]=value
            self.assertEqual(self.post(data).status_code,400,(key,value))
        self.assertEqual(self.client.get('/api/records').json,[])

    def test_csrf_and_private_history(self):
        self.assertEqual(self.client.post('/api/records',json=self.data()).status_code,403)
        self.post(self.data())
        other = app.test_client()
        self.assertEqual(other.get('/api/records').json,[])
        workbook = load_workbook(BytesIO(other.get('/download').data))
        self.assertEqual(workbook.active.max_row,1)

    def test_unconfigured_and_partial_failure(self):
        result=self.post(self.data()).json
        self.assertEqual(result['notifications'],{'father':'not_configured','teacher':'not_configured'})
        with patch('app.send_whatsapp',side_effect=['failed','accepted']):
            result=self.post(self.data()).json
        self.assertEqual(result['notifications'],{'father':'failed','teacher':'accepted'})

    def test_spreadsheet_strings_are_not_formulas(self):
        data=self.data();data['student_name']='=1+1';self.post(data)
        book=load_workbook(BytesIO(self.client.get('/download').data))
        self.assertEqual(book.active['B2'].data_type,'s')

    def test_unworked_subjects_and_empty_day(self):
        data=self.data()
        for key,_,_ in SUBJECTS:
            for suffix in ('correct','wrong','blank'):
                data[key+'_'+suffix]=''
        self.assertEqual(self.post(data).status_code,400)
        data['matematik_correct']='125'
        response=self.post(data)
        self.assertEqual(response.status_code,201)
        self.assertEqual(response.json['total_questions'],125)
        self.assertEqual(response.json['subjects'][0]['total'],0)

    def test_old_exams_not_relabelled_and_same_day_entries(self):
        self.post(self.data())
        self.post(self.data())
        with self.client.session_transaction() as session:
            owner=session['owner']
        with db() as connection:
            connection.execute('INSERT INTO records VALUES (?, ?, ?, ?, ?)',
                               ('legacy',owner,'2026-01-01',json.dumps({'exam_name':'Old exam'}),'{}'))
        self.assertEqual(len(self.client.get('/api/records').json),2)
        workbook=load_workbook(BytesIO(self.client.get('/download').data))
        self.assertEqual(workbook.active.max_row,3)
        self.assertEqual(workbook.active['C2'].value,330)

    def test_pwa_and_routes(self):
        for path in ['/', '/report', '/qr', '/sw.js', '/static/app.js', '/static/app.css','/static/manifest.webmanifest','/static/icon-192.png','/static/icon-512.png']:
            response=self.client.get(path)
            self.assertEqual(response.status_code,200,path)
            response.close()
        self.assertEqual(self.client.get('/api/records').headers['Cache-Control'],'no-store')

    def test_provider_template_request(self):
        settings=dict(TWILIO_ACCOUNT_SID='ACtest',TWILIO_AUTH_TOKEN='test',TWILIO_WHATSAPP_FROM='+12345678900',TWILIO_DAILY_CONTENT_SID='HXtest')
        record=self.post(self.data()).json
        response=MagicMock();response.__enter__.return_value.read.return_value=b'{"sid":"SMtest"}'
        with patch.dict(os.environ,settings),patch('app.urlopen',return_value=response) as request:
            self.assertEqual(send_whatsapp('+905321234567',record),'accepted')
            body=request.call_args.args[0].data.decode()
            self.assertIn('ContentSid=HXtest',body)
            self.assertIn('To=whatsapp%3A%2B905321234567',body)
            variables=json.loads(parse_qs(body)['ContentVariables'][0])
            self.assertEqual(variables['3'],'330')
            self.assertIn('55 soru (50 D, 3 Y, 2 B)',variables['4'])


if __name__ == '__main__':
    unittest.main()
