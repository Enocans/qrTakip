import os
from datetime import date
from io import BytesIO

import qrcode
from flask import Flask, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
if os.environ.get("VERCEL") == "1" or os.environ.get("VERCEL_ENV"):
    DATA_DIR = "/tmp/qrTakip_data"
else:
    DATA_DIR = os.path.join(BASE_DIR, "data")

os.makedirs(DATA_DIR, exist_ok=True)
EXCEL_PATH = os.path.join(DATA_DIR, "student_tracker.xlsx")
DEFAULT_TEACHER = "Koç Abdulaziz"

app = Flask(__name__)

SUBJECTS = [
    ("turkce", "Türkçe"),
    ("matematik", "Matematik"),
    ("fen", "Fen"),
    ("inkilap", "İnkılap"),
    ("ingilizce", "İngilizce"),
    ("din", "Din"),
]

TARGETS = {
    "turkce": 100,
    "matematik": 100,
    "fen": 75,
    "inkilap": 50,
    "ingilizce": 40,
    "din": 20,
}

HEADERS = [
    "Tarih",
    "Öğrenci Adı",
    "Baba Adı",
    "Öğretmen",
    "Telefon",
    "Türkçe Doğru",
    "Türkçe Yanlış",
    "Türkçe Net",
    "Matematik Doğru",
    "Matematik Yanlış",
    "Matematik Net",
    "Fen Doğru",
    "Fen Yanlış",
    "Fen Net",
    "İnkılap Doğru",
    "İnkılap Yanlış",
    "İnkılap Net",
    "İngilizce Doğru",
    "İngilizce Yanlış",
    "İngilizce Net",
    "Din Doğru",
    "Din Yanlış",
    "Din Net",
    "Toplam Net",
    "Toplam Soru",
    "Çalışma Süresi",
    "Durum",
    "Not",
]


def calculate_net(correct, wrong):
    return correct - (wrong / 4)


def ensure_excel_file():
    if not os.path.exists(EXCEL_PATH):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "LGS Takip"
        worksheet.append(HEADERS)
        workbook.save(EXCEL_PATH)
        return

    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook.active
    if worksheet.max_row == 1 and worksheet["A1"].value is None:
        worksheet.append(HEADERS)
        workbook.save(EXCEL_PATH)


def save_record(form_data):
    ensure_excel_file()
    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook.active

    subject_totals = {}
    total_correct = 0
    total_wrong = 0

    for key, label in SUBJECTS:
        correct = int(form_data.get(f"{key}_correct", 0) or 0)
        wrong = int(form_data.get(f"{key}_wrong", 0) or 0)
        net = calculate_net(correct, wrong)
        subject_totals[key] = {
            "correct": correct,
            "wrong": wrong,
            "net": net,
            "label": label,
        }
        total_correct += correct
        total_wrong += wrong

    total_net = total_correct - (total_wrong / 4)
    total_questions = sum(
        int(form_data.get(f"{key}_correct", 0) or 0) + int(form_data.get(f"{key}_wrong", 0) or 0)
        for key, _ in SUBJECTS
    )

    date_value = form_data.get("date") or date.today().strftime("%d.%m.%Y")
    status = "Hedefe Uygun" if total_net >= 120 else "İzlenmeli"

    row = [
        date_value,
        form_data.get("student_name", ""),
        form_data.get("father_name", ""),
        form_data.get("teacher_name", ""),
        form_data.get("phone", ""),
        subject_totals["turkce"]["correct"],
        subject_totals["turkce"]["wrong"],
        subject_totals["turkce"]["net"],
        subject_totals["matematik"]["correct"],
        subject_totals["matematik"]["wrong"],
        subject_totals["matematik"]["net"],
        subject_totals["fen"]["correct"],
        subject_totals["fen"]["wrong"],
        subject_totals["fen"]["net"],
        subject_totals["inkilap"]["correct"],
        subject_totals["inkilap"]["wrong"],
        subject_totals["inkilap"]["net"],
        subject_totals["ingilizce"]["correct"],
        subject_totals["ingilizce"]["wrong"],
        subject_totals["ingilizce"]["net"],
        subject_totals["din"]["correct"],
        subject_totals["din"]["wrong"],
        subject_totals["din"]["net"],
        round(total_net, 2),
        total_questions,
        form_data.get("study_time", ""),
        status,
        form_data.get("notes", ""),
    ]

    worksheet.append(row)
    workbook.save(EXCEL_PATH)


@app.route("/")
def index():
    today = date.today().strftime("%d.%m.%Y")
    return render_template("index.html", today=today, teacher_name=DEFAULT_TEACHER)


@app.route("/submit", methods=["POST"])
def submit_record():
    save_record(request.form)
    return redirect(url_for("report"))


@app.route("/panel")
@app.route("/report")
def report():
    ensure_excel_file()
    workbook = load_workbook(EXCEL_PATH)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    headers = rows[0] if rows else []
    data_rows = rows[1:] if len(rows) > 1 else []
    return render_template("report.html", headers=headers, data_rows=data_rows)


@app.route("/download")
def download_excel():
    ensure_excel_file()
    return send_file(EXCEL_PATH, as_attachment=True, download_name="student_tracker.xlsx")


@app.route("/qr")
def qr_code():
    base_url = os.environ.get("PUBLIC_URL") or request.host_url.rstrip("/")
    if not base_url.startswith(("http://", "https://")):
        base_url = f"https://{base_url}"

    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(base_url)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return send_file(buffer, mimetype="image/png")


if __name__ == "__main__":
    ensure_excel_file()
    port = int(os.environ.get("PORT", 5000))
    app.run(debug=False, host="0.0.0.0", port=port)
